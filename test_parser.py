import openpyxl
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Ngan hang 2026.xlsm', data_only=True)

def parse_month_sheet(ws):
    units_data = {}
    current_unit = None
    unit_rows = []
    
    for r in range(5, ws.max_row + 1):
        col_a = ws.cell(r, 1).value
        col_b = ws.cell(r, 2).value
        if col_a and str(col_a).strip():
            current_unit = str(col_a).strip()
            
        if not col_b:
            continue
            
        bank_name = str(col_b).strip()
        
        def safe_float(v):
            if isinstance(v, (int, float)):
                return float(v)
            try:
                return float(str(v).replace(',', ''))
            except Exception:
                return 0.0

        hd_unt = safe_float(ws.cell(r, 3).value)
        st_unt = safe_float(ws.cell(r, 4).value)
        hd_unc = safe_float(ws.cell(r, 5).value)
        st_unc = safe_float(ws.cell(r, 6).value)
        hd_app = safe_float(ws.cell(r, 7).value)
        st_app = safe_float(ws.cell(r, 8).value)
        hd_quay = safe_float(ws.cell(r, 9).value)
        st_quay = safe_float(ws.cell(r, 10).value)
        hd_tong = safe_float(ws.cell(r, 11).value)
        st_tong = safe_float(ws.cell(r, 12).value)
        
        row_dict = {
            'unit': current_unit,
            'bank': bank_name,
            'hd_unt': hd_unt,
            'st_unt': st_unt,
            'hd_unc': hd_unc,
            'st_unc': st_unc,
            'hd_app': hd_app,
            'st_app': st_app,
            'hd_quay': hd_quay,
            'st_quay': st_quay,
            'hd_tong': hd_tong,
            'st_tong': st_tong,
        }
        
        if 'TỔNG CỘNG' in bank_name.upper():
            if current_unit not in units_data:
                units_data[current_unit] = {'banks': unit_rows, 'total': row_dict}
            else:
                units_data[current_unit]['total'] = row_dict
            unit_rows = []
        else:
            unit_rows.append(row_dict)
            
    return units_data

print('Testing parse Thang 09:')
data_09 = parse_month_sheet(wb['Thang 09'])
print('Units parsed:', list(data_09.keys()))
for u, d in list(data_09.items()):
    tot = d.get('total', {})
    b_len = len(d.get('banks', []))
    print(f"Unit: {u:15} | Banks: {b_len:2} | Total HD: {tot.get('hd_tong', 0):>10,.0f} | Total Tien: {tot.get('st_tong', 0):>18,.0f}")
