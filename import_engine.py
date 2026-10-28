import openpyxl
import os
import sys
import json
import re

from data_engine import (
    load_all_data, 
    normalize_unit_name, 
    get_short_bank_name, 
    classify_bank, 
    safe_float, 
    CACHE_JSON_PATH,
    DATA_EXCEL_PATH
)

# Ensure UTF-8 output
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def parse_raw_sql_file(file_path, month_name):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    # Read headers from row 1
    headers = [str(ws.cell(1, c).value or '').strip() for c in range(1, ws.max_column + 1)]
    
    # Find column indices
    col_map = {}
    for idx, h in enumerate(headers, 1):
        h_u = h.upper()
        if 'TEN_DVIQLY' in h_u: col_map['unit'] = idx
        elif 'MA_DVIQLY' in h_u and 'unit_code' not in col_map: col_map['unit_code'] = idx
        elif 'TEN_TNGAN' in h_u: col_map['bank'] = idx
        elif 'SO_HDON_UT' in h_u: col_map['hd_unt'] = idx
        elif 'SO_TIEN_UT' in h_u: col_map['st_unt'] = idx
        elif 'SO_HDON_UC' in h_u: col_map['hd_unc'] = idx
        elif 'SO_TIEN_UC' in h_u: col_map['st_unc'] = idx
        elif 'SO_HDON_DT_TD' in h_u: col_map['hd_app'] = idx
        elif 'SO_TIEN_DT_TD' in h_u: col_map['st_app'] = idx
        elif 'SO_HDON_CK' in h_u: col_map['hd_quay'] = idx
        elif 'SO_TIEN_CK' in h_u: col_map['st_quay'] = idx

    if 'unit' not in col_map or 'bank' not in col_map:
        raise ValueError("File Excel không đúng định dạng trích xuất SQL CMIS (thiếu cột TEN_DVIQLY hoặc TEN_TNGAN)")

    units_data = {}
    for r in range(2, ws.max_row + 1):
        u_raw = ws.cell(r, col_map['unit']).value
        b_raw = ws.cell(r, col_map['bank']).value
        if not u_raw or not b_raw:
            continue
            
        unit_norm = normalize_unit_name(u_raw)
        bank_name = str(b_raw).strip()
        
        hd_unt = safe_float(ws.cell(r, col_map.get('hd_unt', 0)).value) if 'hd_unt' in col_map else 0.0
        st_unt = safe_float(ws.cell(r, col_map.get('st_unt', 0)).value) if 'st_unt' in col_map else 0.0
        hd_unc = safe_float(ws.cell(r, col_map.get('hd_unc', 0)).value) if 'hd_unc' in col_map else 0.0
        st_unc = safe_float(ws.cell(r, col_map.get('st_unc', 0)).value) if 'st_unc' in col_map else 0.0
        hd_app = safe_float(ws.cell(r, col_map.get('hd_app', 0)).value) if 'hd_app' in col_map else 0.0
        st_app = safe_float(ws.cell(r, col_map.get('st_app', 0)).value) if 'st_app' in col_map else 0.0
        hd_quay = safe_float(ws.cell(r, col_map.get('hd_quay', 0)).value) if 'hd_quay' in col_map else 0.0
        st_quay = safe_float(ws.cell(r, col_map.get('st_quay', 0)).value) if 'st_quay' in col_map else 0.0
        
        hd_tong = hd_unt + hd_unc + hd_app + hd_quay
        st_tong = st_unt + st_unc + st_app + st_quay
        
        if unit_norm not in units_data:
            units_data[unit_norm] = {
                'unit_name': unit_norm,
                'banks': [],
                'total': {
                    'unit': unit_norm,
                    'bank': 'TỔNG CỘNG',
                    'bank_short': 'TỔNG CỘNG',
                    'category': 'Tổng hợp',
                    'hd_unt': 0.0, 'st_unt': 0.0,
                    'hd_unc': 0.0, 'st_unc': 0.0,
                    'hd_app': 0.0, 'st_app': 0.0,
                    'hd_quay': 0.0, 'st_quay': 0.0,
                    'hd_tong': 0.0, 'st_tong': 0.0
                }
            }
            
        short_b = get_short_bank_name(bank_name)
        cat = classify_bank(bank_name)
        
        b_entry = {
            'unit': unit_norm,
            'bank': bank_name,
            'bank_short': short_b,
            'category': cat,
            'hd_unt': hd_unt, 'st_unt': st_unt,
            'hd_unc': hd_unc, 'st_unc': st_unc,
            'hd_app': hd_app, 'st_app': st_app,
            'hd_quay': hd_quay, 'st_quay': st_quay,
            'hd_tong': hd_tong, 'st_tong': st_tong
        }
        units_data[unit_norm]['banks'].append(b_entry)
        
        tot = units_data[unit_norm]['total']
        tot['hd_unt'] += hd_unt
        tot['st_unt'] += st_unt
        tot['hd_unc'] += hd_unc
        tot['st_unc'] += st_unc
        tot['hd_app'] += hd_app
        tot['st_app'] += st_app
        tot['hd_quay'] += hd_quay
        tot['st_quay'] += st_quay
        tot['hd_tong'] += hd_tong
        tot['st_tong'] += st_tong

    # Create TOÀN CÔNG TY summary
    company_banks = {}
    company_total = {
        'unit': 'TOÀN CÔNG TY',
        'bank': 'TỔNG CỘNG',
        'bank_short': 'TỔNG CỘNG',
        'category': 'Tổng hợp',
        'hd_unt': 0.0, 'st_unt': 0.0,
        'hd_unc': 0.0, 'st_unc': 0.0,
        'hd_app': 0.0, 'st_app': 0.0,
        'hd_quay': 0.0, 'st_quay': 0.0,
        'hd_tong': 0.0, 'st_tong': 0.0
    }
    
    for u_k, u_v in units_data.items():
        t = u_v['total']
        for k in ['hd_unt', 'st_unt', 'hd_unc', 'st_unc', 'hd_app', 'st_app', 'hd_quay', 'st_quay', 'hd_tong', 'st_tong']:
            company_total[k] += t[k]
            
        for b in u_v['banks']:
            bn = b['bank']
            if bn not in company_banks:
                company_banks[bn] = {
                    'unit': 'TOÀN CÔNG TY',
                    'bank': bn,
                    'bank_short': b['bank_short'],
                    'category': b['category'],
                    'hd_unt': 0.0, 'st_unt': 0.0,
                    'hd_unc': 0.0, 'st_unc': 0.0,
                    'hd_app': 0.0, 'st_app': 0.0,
                    'hd_quay': 0.0, 'st_quay': 0.0,
                    'hd_tong': 0.0, 'st_tong': 0.0
                }
            for k in ['hd_unt', 'st_unt', 'hd_unc', 'st_unc', 'hd_app', 'st_app', 'hd_quay', 'st_quay', 'hd_tong', 'st_tong']:
                company_banks[bn][k] += b[k]

    company_banks_list = list(company_banks.values())
    company_banks_list.sort(key=lambda x: x['st_tong'], reverse=True)
    
    units_data['TOÀN CÔNG TY'] = {
        'unit_name': 'TOÀN CÔNG TY',
        'banks': company_banks_list,
        'total': company_total
    }
    
    # Save into data cache
    all_data = load_all_data()
    all_data['months_data'][month_name] = units_data
    if month_name not in all_data['available_months']:
        # insert before Luy ke if exists
        if "Lũy kế 9 tháng" in all_data['available_months']:
            idx = all_data['available_months'].index("Lũy kế 9 tháng")
            all_data['available_months'].insert(idx, month_name)
        else:
            all_data['available_months'].append(month_name)
            
    with open(CACHE_JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(all_data, f, ensure_ascii=False, indent=2)

    return all_data

if __name__ == "__main__":
    # Test importing Thang 092026.xlsx
    test_sql_file = r"Ngan hang\Du lieu xuat tu doan code\Thang 092026.xlsx"
    if os.path.exists(test_sql_file):
        res = parse_raw_sql_file(test_sql_file, "Thang 09 (SQL)")
        print("Imported SQL successfully. Total months:", len(res['available_months']))
