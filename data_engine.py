import openpyxl
import os
import json
import re
import sys

# Ensure UTF-8 output
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_EXCEL_PATH = os.path.join(BASE_DIR, "Ngan hang 2026.xlsm")
CACHE_JSON_PATH = os.path.join(BASE_DIR, "data_cache.json")

UNIT_ORDER = [
    "ĐÀ LẠT",
    "BẢO LỘC",
    "ĐƠN DƯƠNG",
    "DI LINH",
    "ĐỨC TRỌNG",
    "LÂM HÀ",
    "ĐẠ HUOAI",
    "ĐẠ TẺH",
    "CÁT TIÊN",
    "BẢO LÂM",
    "LANGBIANG",
    "ĐAM RÔNG"
]

def safe_float(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        s = str(v).replace(',', '').replace(' ', '').strip()
        return float(s) if s else 0.0
    except Exception:
        return 0.0

def normalize_unit_name(name):
    if not name:
        return ""
    name_clean = str(name).strip().upper()
    name_clean = re.sub(r'^(ĐIỆN LỰC\s+|DL\s+)', '', name_clean)
    name_clean = name_clean.replace("LANG BIANG", "LANGBIANG")
    return name_clean

def get_short_bank_name(name):
    if not name:
        return ""
    n = str(name).upper()
    if "VIETCOMBANK" in n or "NGOẠI THƯƠNG" in n:
        return "VietcomBank"
    if "VIETINBANK" in n or "CÔNG THƯƠNG" in n:
        return "VietinBank"
    if "AGRIBANK" in n or "NÔNG NGHIỆP" in n:
        return "Agribank"
    if "BIDV" in n or "ĐẦU TƯ VÀ PHÁT TRIỂN" in n:
        return "BIDV"
    if "MB" in n or "QUÂN ĐỘI" in n and "VIETTEL" not in n:
        return "MB Bank"
    if "ACB" in n or "Á CHÂU" in n:
        return "ACB"
    if "SACCOM" in n or "SACOMBANK" in n or "SÀI GÒN THƯƠNG TÍN" in n:
        return "Sacombank"
    if "VIB" in n or "QUỐC TẾ" in n:
        return "VIB"
    if "HDBANK" in n or "PHÁT TRIỂN TP" in n:
        return "HDBank"
    if "PVCOMBANK" in n or "ĐẠI CHÚNG" in n:
        return "PVCombank"
    if "OCB" in n or "PHƯƠNG ĐÔNG" in n:
        return "OCB"
    if "ABBANK" in n or "AN BÌNH" in n:
        return "ABBank"
    if "MOMO" in n or "M_SERVICE" in n:
        return "MoMo"
    if "VNPAY" in n:
        return "VNPay"
    if "ZALOPAY" in n or "ZALO PAY" in n or "ZION" in n:
        return "ZaloPay"
    if "PAYOO" in n or "VIETUNION" in n:
        return "Payoo"
    if "PAYTECH" in n:
        return "Paytech"
    if "SHOPEEPAY" in n:
        return "ShopeePay"
    if "VIETTEL" in n:
        return "Viettel Lâm Đồng"
    if "VNPT" in n:
        return "VNPT Media"
    if "VIETNAM POST" in n or "BƯU ĐIỆN" in n:
        return "Vietnam Post"
    if "FINVIET" in n:
        return "FinViet"
    return name[:30]

def classify_bank(bank_name):
    short = get_short_bank_name(bank_name)
    if short in ["Agribank", "VietcomBank", "VietinBank", "BIDV"]:
        return "Big 4 Ngân hàng"
    if short in ["MoMo", "VNPay", "ZaloPay", "Payoo", "Paytech", "ShopeePay", "Viettel Lâm Đồng", "VNPT Media", "Vietnam Post", "FinViet"]:
        return "Ví điện tử & Fintech"
    return "Ngân hàng TMCP khác"

def parse_month_sheet(ws):
    units_data = {}
    current_unit = None
    unit_rows = []
    
    max_r = ws.max_row
    for r in range(5, max_r + 1):
        col_a = ws.cell(r, 1).value
        col_b = ws.cell(r, 2).value
        
        if col_a and str(col_a).strip():
            raw_unit = str(col_a).strip()
            norm = normalize_unit_name(raw_unit)
            if norm:
                current_unit = norm
                
        if not col_b:
            continue
            
        bank_name = str(col_b).strip()
        if not bank_name:
            continue
            
        hd_unt = safe_float(ws.cell(r, 3).value)
        st_unt = safe_float(ws.cell(r, 4).value)
        hd_unc = safe_float(ws.cell(r, 5).value)
        st_unc = safe_float(ws.cell(r, 6).value)
        hd_app = safe_float(ws.cell(r, 7).value)
        st_app = safe_float(ws.cell(r, 8).value)
        hd_quay = safe_float(ws.cell(r, 9).value)
        st_quay = safe_float(ws.cell(r, 10).value)
        hd_tong = safe_float(ws.cell(r, 11).value)
        st_tong = safe_float(ws.cell(r, 12).value)
        
        if hd_tong == 0 and (hd_unt + hd_unc + hd_app + hd_quay) > 0:
            hd_tong = hd_unt + hd_unc + hd_app + hd_quay
        if st_tong == 0 and (st_unt + st_unc + st_app + st_quay) > 0:
            st_tong = st_unt + st_unc + st_app + st_quay

        short_bank = get_short_bank_name(bank_name)
        category = classify_bank(bank_name)

        row_dict = {
            'unit': current_unit,
            'bank': bank_name,
            'bank_short': short_bank,
            'category': category,
            'hd_unt': hd_unt,
            'st_unt': st_unt,
            'hd_unc': hd_unc,
            'st_unc': st_unc,
            'hd_app': hd_app,
            'st_app': st_app,
            'hd_quay': hd_quay,
            'st_quay': st_quay,
            'hd_tong': hd_tong,
            'st_tong': st_tong
        }
        
        if 'TỔNG CỘNG' in bank_name.upper():
            if current_unit:
                units_data[current_unit] = {
                    'unit_name': current_unit,
                    'banks': unit_rows,
                    'total': row_dict
                }
            unit_rows = []
        else:
            unit_rows.append(row_dict)
            
    # Calculate company summary (sum of all 12 units)
    company_banks = {}
    company_total = {
        'unit': 'TOÀN CÔNG TY',
        'bank': 'TỔNG CỘNG',
        'bank_short': 'TỔNG CỘNG',
        'category': 'Tổng hợp',
        'hd_unt': 0.0, 'st_unt': 0.0,
        'hd_unc': 0.0, 'st_unc': 0.0,
        'hd_app': 0.0, 'st_app': 0.0,
        'hd_quay': 0.0, 'st_quay': 0.0,
        'hd_tong': 0.0, 'st_tong': 0.0
    }
    
    for u_key, u_val in units_data.items():
        if u_key in ['TRỤ SỞ CHÍNH', 'TOÀN CÔNG TY']:
            continue
        tot = u_val['total']
        for k in ['hd_unt', 'st_unt', 'hd_unc', 'st_unc', 'hd_app', 'st_app', 'hd_quay', 'st_quay', 'hd_tong', 'st_tong']:
            company_total[k] += tot[k]
            
        for b in u_val['banks']:
            b_name = b['bank']
            if b_name not in company_banks:
                company_banks[b_name] = {
                    'unit': 'TOÀN CÔNG TY',
                    'bank': b_name,
                    'bank_short': b['bank_short'],
                    'category': b['category'],
                    'hd_unt': 0.0, 'st_unt': 0.0,
                    'hd_unc': 0.0, 'st_unc': 0.0,
                    'hd_app': 0.0, 'st_app': 0.0,
                    'hd_quay': 0.0, 'st_quay': 0.0,
                    'hd_tong': 0.0, 'st_tong': 0.0
                }
            for k in ['hd_unt', 'st_unt', 'hd_unc', 'st_unc', 'hd_app', 'st_app', 'hd_quay', 'st_quay', 'hd_tong', 'st_tong']:
                company_banks[b_name][k] += b[k]
                
    company_banks_list = list(company_banks.values())
    company_banks_list.sort(key=lambda x: x['st_tong'], reverse=True)
    
    units_data['TOÀN CÔNG TY'] = {
        'unit_name': 'TOÀN CÔNG TY',
        'banks': company_banks_list,
        'total': company_total
    }
    
    return units_data

def load_all_data(excel_path=DATA_EXCEL_PATH, force_reload=False):
    if not force_reload and os.path.exists(CACHE_JSON_PATH):
        try:
            with open(CACHE_JSON_PATH, 'r', encoding='utf-8') as f:
                cached = json.load(f)
                return cached
        except Exception:
            pass

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    months_data = {}
    available_months = []
    
    # Check Thang 01 to Thang 12
    for m in range(1, 13):
        sheet_name = f"Thang {m:02d}"
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            m_data = parse_month_sheet(ws)
            if m_data and 'TOÀN CÔNG TY' in m_data and m_data['TOÀN CÔNG TY']['total']['st_tong'] > 0:
                months_data[sheet_name] = m_data
                available_months.append(sheet_name)
                
    # Parse Luy ke
    if "Luy ke" in wb.sheetnames:
        luy_ke_data = parse_month_sheet(wb["Luy ke"])
        if luy_ke_data and 'TOÀN CÔNG TY' in luy_ke_data and luy_ke_data['TOÀN CÔNG TY']['total']['st_tong'] > 0:
            months_data["Lũy kế 9 tháng"] = luy_ke_data
            available_months.append("Lũy kế 9 tháng")

    distinct_banks = set()
    distinct_categories = set()
    if available_months:
        first_m = available_months[0]
        for b in months_data[first_m]['TOÀN CÔNG TY']['banks']:
            distinct_banks.add(b['bank_short'])
            distinct_categories.add(b['category'])

    result = {
        'available_months': available_months,
        'units': UNIT_ORDER,
        'banks': sorted(list(distinct_banks)),
        'categories': sorted(list(distinct_categories)),
        'months_data': months_data
    }
    
    try:
        with open(CACHE_JSON_PATH, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Error saving cache: {e}")
        
    return result

if __name__ == "__main__":
    data = load_all_data(force_reload=True)
    print("Parsed months:", len(data['available_months']))
    print("Units:", len(data['units']))
    print("Banks:", len(data['banks']))
    print("SUCCESS")
